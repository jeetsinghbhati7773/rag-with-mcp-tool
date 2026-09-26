from __future__ import annotations

import io
import logging
from typing import List, Dict, Any
from langchain_core.documents import Document

logger = logging.getLogger(__name__)


def load_pdf_bytes(file_bytes: bytes, filename: str) -> List[Document]:
    """Extracts text page by page from PDF file bytes."""
    import pypdf
    docs = []
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    total_pages = len(reader.pages)
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": filename,
                        "document_name": filename,
                        "page_number": idx + 1,
                        "total_pages": total_pages,
                    },
                )
            )
    return docs


def load_docx_bytes(file_bytes: bytes, filename: str) -> List[Document]:
    """Extracts text paragraphs from DOCX file bytes."""
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return [
        Document(
            page_content=full_text,
            metadata={
                "source": filename,
                "document_name": filename,
                "page_number": 1,
                "total_pages": 1,
            },
        )
    ]


def load_txt_bytes(file_bytes: bytes, filename: str) -> List[Document]:
    """Extracts plain text / markdown file bytes."""
    text = file_bytes.decode("utf-8", errors="ignore")
    return [
        Document(
            page_content=text,
            metadata={
                "source": filename,
                "document_name": filename,
                "page_number": 1,
                "total_pages": 1,
            },
        )
    ]


def load_document_bytes(file_bytes: bytes, filename: str) -> List[Document]:
    """Dispatches document extraction based on file extension."""
    ext = filename.lower().split(".")[-1]
    if ext == "pdf":
        return load_pdf_bytes(file_bytes, filename)
    elif ext in ("docx", "doc"):
        return load_docx_bytes(file_bytes, filename)
    elif ext in ("txt", "md", "markdown"):
        return load_txt_bytes(file_bytes, filename)
    else:
        logger.warning("Unsupported extension %s; attempting text decode.", ext)
        return load_txt_bytes(file_bytes, filename)
